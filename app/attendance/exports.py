# ---------------------------------------------------------------------------
# app/attendance/exports.py
# Server-side export builders for a single day's attendance:
#   * build_excel(day)  -> .xlsx bytes (openpyxl)
#   * build_png(day)    -> .png bytes (Pillow)
#
# Both reuse day_records() so the data grouping lives in one place and the
# export format never drifts from the CSV.
# ---------------------------------------------------------------------------

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from .services import get_session

# The project's pixel palette (matches static/css/style.css).
PALETTE = {
    "dark_red": "#BE1A1A",     # background
    "red": "#D0311E",          # headings
    "yellow": "#F7D87F",       # absent highlight
    "cream": "#F8EBAB",        # text
}

REFERENCE = {
    "green": "DCFCE7",
    "amber": "FDE68A",
    "yellow": "FFE600",
    "yellow_soft": "FFF2B0",
    "black": "000000",
    "white": "FFFFFF",
}


def group_records(records):
    """Group attendance records by status, sorted by English name."""
    grouped = {"present": [], "late": [], "absent": []}
    for record in sorted(records, key=lambda r: r.student.name_en.lower()):
        grouped.setdefault(record.status, []).append(record)
    return grouped


def day_records(day):
    """Return a day's records grouped by status."""
    session = get_session(day)
    if session is None:
        return {"present": [], "late": [], "absent": []}
    return group_records(session.records)


def _safe_spreadsheet_value(value):
    """Prevent spreadsheet formula injection in text cells."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_excel_section(
    sheet,
    row,
    title,
    records,
    *,
    header_fill,
    row_fill,
    row_font,
    empty_message,
):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    heading = sheet.cell(row=row, column=1, value=title)
    heading.font = Font(bold=True, size=14, color=REFERENCE["black"])
    heading.fill = PatternFill("solid", fgColor=header_fill)
    heading.alignment = Alignment(horizontal="left")
    row += 1

    for column, label in enumerate(("#", "Name (English)", "الاسم (عربي)"), start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = Font(bold=True, size=11)
        cell.fill = PatternFill("solid", fgColor=header_fill)
    sheet.cell(row=row, column=3).alignment = Alignment(
        horizontal="right", readingOrder=2
    )
    row += 1

    if records:
        for record in records:
            values = (
                record.student_id,
                _safe_spreadsheet_value(record.student.name_en),
                _safe_spreadsheet_value(record.student.name_ar),
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = row_font
                cell.fill = PatternFill("solid", fgColor=row_fill)
                if column == 3:
                    cell.alignment = Alignment(horizontal="right", readingOrder=2)
            row += 1
    else:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = sheet.cell(row=row, column=1, value=empty_message)
        cell.font = row_font
        cell.fill = PatternFill("solid", fgColor=row_fill)
        row += 1

    return row + 1


def build_excel(day):
    """Return the reference-style bilingual .xlsx workbook as bytes."""
    records = day_records(day)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.sheet_view.showGridLines = False

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 34

    sheet.merge_cells("A1:C1")
    sheet["A1"] = (
        "Egypt-Japan KOSEN Institute — Class C1 Attendance / "
        "معهد مصر اليابان كوسن — حضور وغياب فصل C1"
    )
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34

    sheet["A2"] = f"Date / التاريخ: {day.isoformat()}"
    sheet["B2"] = (
        f"Present / حاضر: {len(records['present'])}   "
        f"Late / متأخر: {len(records['late'])}"
    )
    sheet["C2"] = f"Absent / غائب: {len(records['absent'])}"
    for coordinate in ("A2", "B2", "C2"):
        sheet[coordinate].font = Font(bold=True, size=12)

    row = 4
    row = _write_excel_section(
        sheet,
        row,
        "Attended / الحاضرون",
        records["present"],
        header_fill=REFERENCE["green"],
        row_fill=REFERENCE["white"],
        row_font=Font(size=11),
        empty_message="No present students / لا يوجد حاضرون",
    )
    row = _write_excel_section(
        sheet,
        row,
        "Late / المتأخرون",
        records["late"],
        header_fill=REFERENCE["amber"],
        row_fill=REFERENCE["yellow_soft"],
        row_font=Font(bold=True, size=12),
        empty_message="Nobody late / لا يوجد متأخرون",
    )
    _write_excel_section(
        sheet,
        row,
        "Absent / الغائبون",
        records["absent"],
        header_fill=REFERENCE["yellow"],
        row_fill=REFERENCE["yellow"],
        row_font=Font(bold=True, size=18),
        empty_message="No absent students — everyone attended! / لا يوجد غياب — الجميع حاضر!",
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _load_font(size):
    """Return a usable font, falling back to Pillow's default.

    We do not bundle a TTF, so on a fresh host this returns the built-in
    bitmap font. That is small but readable and, crucially, never crashes.
    """
    try:
        return ImageFont.truetype("DejaVuSans-Bold.ttf", size)
    except OSError:
        return ImageFont.load_default()


def build_png(day):
    """Return a poster-style .png (bytes) listing late and absent students."""
    records = day_records(day)
    absent = [r.student.name_en for r in records.get("absent", [])]
    late = [r.student.name_en for r in records.get("late", [])]

    width = 900
    line_height = 46
    header_height = 150
    late_block = 60 + len(late) * line_height if late else 0
    absent_block = 60 + max(len(absent), 1) * line_height
    height = header_height + late_block + absent_block + 80

    image = Image.new("RGB", (width, height), PALETTE["dark_red"])
    draw = ImageDraw.Draw(image)

    title_font = _load_font(40)
    small_font = _load_font(22)
    name_font = _load_font(28)

    # Header.
    draw.text((40, 40), "C1 ATTENDANCE", font=title_font, fill=PALETTE["cream"])
    draw.text((40, 100), day.isoformat(), font=small_font, fill=PALETTE["cream"])

    y = header_height

    # Late section.
    if late:
        draw.text((40, y), "LATE", font=small_font, fill=PALETTE["red"])
        y += 40
        for name in late:
            draw.text((60, y), f"- {name}", font=name_font, fill=PALETTE["cream"])
            y += line_height
        y += 20

    # Absent section.
    draw.text((40, y), "ABSENT", font=small_font, fill=PALETTE["yellow"])
    y += 40
    if absent:
        for name in absent:
            draw.text((60, y), f"- {name}", font=name_font, fill=PALETTE["yellow"])
            y += line_height
    else:
        draw.text((60, y), "Nobody absent. Nice!", font=name_font, fill=PALETTE["cream"])

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
