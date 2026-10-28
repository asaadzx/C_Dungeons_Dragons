"""Tests for attendance grouping, export layouts, and download responses."""

import io
from datetime import date

from openpyxl import load_workbook
from PIL import Image

from app.extensions import db
from app.models import Student, User


def _seed(app):
    with app.app_context():
        user = User(username="t", role="teacher")
        user.set_password("x")
        s1 = Student(name_en="Present P", name_ar="ح")
        s2 = Student(name_en="Late L", name_ar="خ")
        s3 = Student(name_en="Absent A", name_ar="ط")
        db.session.add_all([user, s1, s2, s3])
        db.session.commit()
        return user.id, s1.id, s2.id, s3.id


def test_group_records_separates_statuses(app):
    user_id, s1, s2, s3 = _seed(app)
    with app.app_context():
        from app.attendance.exports import group_records
        from app.attendance.services import save_attendance
        from app.models import AttendanceSession

        save_attendance(
            date(2026, 10, 22),
            {s1: "present", s2: "late", s3: "absent"},
            {},
            user_id,
        )
        session = AttendanceSession.query.one()
        groups = group_records(session.records)

        assert [r.student.name_en for r in groups["present"]] == ["Present P"]
        assert [r.student.name_en for r in groups["late"]] == ["Late L"]
        assert [r.student.name_en for r in groups["absent"]] == ["Absent A"]


def test_excel_uses_reference_layout(app):
    user_id, s1, s2, s3 = _seed(app)
    with app.app_context():
        from app.attendance.exports import build_excel
        from app.attendance.services import save_attendance

        save_attendance(
            date(2026, 10, 22),
            {s1: "present", s2: "late", s3: "absent"},
            {},
            user_id,
        )
        data = build_excel(date(2026, 10, 22))
        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook.active
        # Flatten every non-empty cell so assertions do not depend on the exact
        # row/column position of each bilingual heading.
        text = " ".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value
        )

        assert "Egypt-Japan KOSEN Institute" in text
        assert "حضور وغياب فصل C1" in text
        assert "2026-10-22" in text
        assert "Attended / الحاضرون" in text
        assert "Late / المتأخرون" in text
        assert "Absent / الغائبون" in text
        assert "Present P" in text
        assert "Late L" in text
        assert "Absent A" in text
        assert "ط" in text
        assert "ح" in text
        assert "خ" in text


def test_excel_shows_all_present_message(app):
    user_id, s1, _s2, _s3 = _seed(app)
    with app.app_context():
        from app.attendance.exports import build_excel
        from app.attendance.services import save_attendance

        save_attendance(date(2026, 10, 22), {s1: "present"}, {}, user_id)
        data = build_excel(date(2026, 10, 22))
        workbook = load_workbook(io.BytesIO(data))
        text = " ".join(
            str(cell.value)
            for row in workbook.active.iter_rows()
            for cell in row
            if cell.value
        )

        assert "No absent students" in text
        assert "لا يوجد غياب" in text


def test_png_uses_reference_white_palette(app):
    user_id, s1, s2, s3 = _seed(app)
    with app.app_context():
        from app.attendance.exports import build_png
        from app.attendance.services import save_attendance

        save_attendance(
            date(2026, 10, 22),
            {s1: "present", s2: "late", s3: "absent"},
            {},
            user_id,
        )
        data = build_png(date(2026, 10, 22))

    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    # The reference design is white; the old export used the site's dark red.
    image = Image.open(io.BytesIO(data)).convert("RGB")
    assert image.getpixel((0, 0)) == (255, 255, 255)
    assert image.getpixel((0, 0)) != (190, 26, 26)


def test_poster_model_has_late_and_absent(app):
    user_id, s1, s2, s3 = _seed(app)
    with app.app_context():
        from app.attendance.exports import poster_model
        from app.attendance.services import save_attendance

        save_attendance(
            date(2026, 10, 22),
            {s1: "present", s2: "late", s3: "absent"},
            {},
            user_id,
        )
        model = poster_model(date(2026, 10, 22))

        assert [r.student.name_en for r in model["late"]] == ["Late L"]
        assert [r.student.name_en for r in model["absent"]] == ["Absent A"]
        assert model["late_count"] == 1
        assert model["absent_count"] == 1


def test_export_routes_require_login(client):
    assert client.get("/attendance/export.xlsx").status_code == 302
    assert client.get("/attendance/export.png").status_code == 302


def test_exports_are_not_cached(client, app):
    """Exports must never be cached: the same URL is re-used after each save.

    Without no-store the browser serves the previous download and the teacher
    sees stale attendance (the "image doesn't update" bug).
    """
    _seed(app)
    _login(client)

    for path in (
        "/attendance/export.csv",
        "/attendance/export.xlsx",
        "/attendance/export.png",
    ):
        response = client.get(path)
        cache_control = response.headers.get("Cache-Control", "")
        assert "no-store" in cache_control, f"{path} is cacheable: {cache_control!r}"


def test_daily_page_has_print_poster(client, app):
    user_id, s1, s2, s3 = _seed(app)
    with app.app_context():
        from app.attendance.services import save_attendance

        save_attendance(
            date(2026, 10, 22),
            {s1: "present", s2: "late", s3: "absent"},
            {},
            user_id,
        )
    _login(client)

    page = client.get("/attendance?date=2026-10-22").data.decode()
    assert 'id="attendance-print"' in page
    assert "attendance-poster" in page
    assert "Late L" in page
    assert "Absent A" in page
    assert "المتأخرون" in page
    assert "الغائبون" in page


def test_export_filenames_match_reference(app, client):
    _seed(app)
    _login(client)

    # PNG/XLSX adopt the reference "c1Class" stem. CSV deliberately keeps its
    # existing lowercase name so external weekly workflows are unaffected.
    for path, expected in (
        ("/attendance/export.png?date=2026-10-22", "2026-10-22-c1Class-attendance.png"),
        ("/attendance/export.xlsx?date=2026-10-22", "2026-10-22-c1Class-attendance.xlsx"),
        ("/attendance/export.csv?date=2026-10-22", "2026-10-22-c1-attendance.csv"),
    ):
        response = client.get(path)
        disposition = response.headers.get("Content-Disposition", "")
        assert f"filename={expected}" in disposition


def _login(client):
    import re

    page = client.get("/auth/login").data.decode()
    token = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
    client.post(
        "/auth/login",
        data={"username": "t", "password": "x", "csrf_token": token},
    )
